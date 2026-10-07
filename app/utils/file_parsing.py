from __future__ import annotations

import csv
import zipfile
from io import BytesIO, StringIO
from pathlib import Path
from typing import Optional

import pandas as pd

from app.utils.column_mapping import map_columns, required_missing
from app.utils.errors import AppError

XLSX_MAGIC = b"PK\x03\x04"
ALLOWED_EXT = {".csv", ".xlsx"}
CSV_DELIMITERS = ",;\t"


def extension_ok(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_EXT


def sniff_kind(path: Path, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise AppError(
            400,
            "INVALID_FILE_FORMAT",
            "Invalid file format. Only .csv and .xlsx supported.",
        )
    raw = path.read_bytes()[:8]
    if ext == ".xlsx":
        if not raw.startswith(XLSX_MAGIC):
            raise AppError(400, "INVALID_FILE_FORMAT", "Invalid file format. Only .csv and .xlsx supported.")
        _guard_xlsx_zip(path)
        return "xlsx"
    if raw.startswith(XLSX_MAGIC):
        raise AppError(400, "INVALID_FILE_FORMAT", "Invalid file format. Only .csv and .xlsx supported.")
    try:
        _decode_sample(path)
    except UnicodeDecodeError as exc:
        raise AppError(400, "UNREADABLE_FILE", "The file could not be decoded as text.") from exc
    return "csv"


def _guard_xlsx_zip(path: Path, uncompressed_cap: int = 524_288_000) -> None:
    try:
        with zipfile.ZipFile(path) as zf:
            total = 0
            for info in zf.infolist():
                total += max(0, int(info.file_size))
                if total > uncompressed_cap:
                    raise AppError(400, "UNREADABLE_FILE", "The spreadsheet is unreadable or exceeds size limits.")
    except AppError:
        raise
    except zipfile.BadZipFile as exc:
        raise AppError(400, "UNREADABLE_FILE", "The spreadsheet could not be read.") from exc


def _decode_sample(path: Path, nbytes: int = 8192) -> tuple[str, str]:
    data = path.read_bytes()[:nbytes]
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"


def peek_headers(path: Path, kind: str, sheet: Optional[str] = None) -> list[str]:
    try:
        if kind == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
                row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
                if not row:
                    raise AppError(400, "EMPTY_FILE", "The file is empty.")
                return [str(c) if c is not None else "" for c in row]
            finally:
                wb.close()
        sample, _enc = _decode_sample(path)
        if not sample.strip():
            raise AppError(400, "EMPTY_FILE", "The file is empty.")
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS)
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(StringIO(sample), dialect)
        header = next(reader, None)
        if not header:
            raise AppError(400, "EMPTY_FILE", "The file is empty.")
        return header
    except AppError:
        raise
    except Exception as exc:
        raise AppError(400, "UNREADABLE_FILE", "The file could not be read.") from exc


def validate_headers(headers: list[str], require_date: bool) -> dict[str, Optional[str]]:
    mapping = map_columns(headers)
    missing = required_missing(mapping, require_date)
    if missing:
        raise AppError(
            422,
            "MISSING_COLUMN",
            "Required column(s) are missing.",
            {"missing": missing, "found": [h for h in headers if str(h).strip()]},
        )
    return mapping


def read_dataframe(path: Path, kind: str, sheet: Optional[str] = None) -> pd.DataFrame:
    try:
        if kind == "xlsx":
            return pd.read_excel(path, sheet_name=sheet if sheet else 0, engine="openpyxl")
        sample, enc = _decode_sample(path, nbytes=16384)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS)
            sep = dialect.delimiter
        except csv.Error:
            sep = ","
        return pd.read_csv(path, encoding=enc, sep=sep)
    except AppError:
        raise
    except Exception as exc:
        raise AppError(400, "UNREADABLE_FILE", "The file could not be read.") from exc


def csv_bytes_from_path(path: Path) -> bytes:
    return path.read_bytes()


def peek_from_bytes_csv(data: bytes) -> list[str]:
    text = data.decode("utf-8-sig")
    reader = csv.reader(StringIO(text))
    return next(reader)
